# SPDX-FileCopyrightText: 2026 Xoe-NovAi
#
# SPDX-License-Identifier: Apache-2.0

"""Tests for Universal Document Reader."""
# AP: AP-OMEGA-DOC-READER-TESTS-v1.0.0
# ⬡ OMEGA ⬡ DOC_READER_TESTS ⬡ 2026-07-13

import pytest
import tempfile
import os
from pathlib import Path

from omega.doc_reader import (
    read_document,
    read_document_with_metadata,
    DocumentReader,
    DocumentMetadata,
)


class TestDocumentReader:
    """Test DocumentReader class."""

    def test_read_docx(self):
        """Test reading .docx file."""
        # Create test docx
        import docx
        doc = docx.Document()
        doc.add_heading("Test Document", 0)
        doc.add_paragraph("This is a test paragraph.")
        doc.add_paragraph("Another paragraph with **bold** text.")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            try:
                text = read_document(f.name)
                assert "Test Document" in text
                assert "test paragraph" in text
                assert "bold" in text
            finally:
                os.unlink(f.name)

    def test_read_pdf(self):
        """Test reading .pdf file."""
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Test PDF Document\nLine 2\nLine 3")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            doc.save(f.name)
            try:
                text = read_document(f.name)
                assert "Test PDF Document" in text
                assert "Line 2" in text
            finally:
                os.unlink(f.name)

    def test_read_odt(self):
        """Test reading .odt file."""
        from odf.opendocument import OpenDocumentText
        from odf.text import P
        doc = OpenDocumentText()
        doc.text.addElement(P(text="Test ODT Document"))
        doc.text.addElement(P(text="Second paragraph"))

        with tempfile.NamedTemporaryFile(suffix=".odt", delete=False) as f:
            doc.save(f.name)
            try:
                text = read_document(f.name)
                assert "Test ODT Document" in text
                assert "Second paragraph" in text
            finally:
                os.unlink(f.name)

    def test_read_rtf(self):
        """Test reading .rtf file."""
        rtf_content = r"{\rtf1\ansi\deff0 Test RTF Document\par Second line\par}"
        with tempfile.NamedTemporaryFile(suffix=".rtf", delete=False, mode='w') as f:
            f.write(rtf_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert "Test RTF Document" in text
            assert "Second line" in text
        finally:
            os.unlink(fname)

    def test_read_html(self):
        """Test reading .html file."""
        html_content = """<html><body>
<h1>Test HTML</h1>
<p>This is a <strong>test</strong> paragraph.</p>
<ul><li>Item 1</li><li>Item 2</li></ul>
</body></html>"""
        with tempfile.NamedTemporaryFile(suffix=".html", delete=False, mode='w') as f:
            f.write(html_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert "Test HTML" in text
            assert "test" in text
            assert "Item 1" in text
        finally:
            os.unlink(fname)

    def test_read_markdown(self):
        """Test reading .md file."""
        md_content = """# Test Markdown

This is **bold** and *italic* text.

- List item 1
- List item 2
"""
        with tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode='w') as f:
            f.write(md_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert "Test Markdown" in text
            assert "bold" in text
            assert "List item 1" in text
        finally:
            os.unlink(fname)

    def test_read_text(self):
        """Test reading .txt file."""
        txt_content = "Plain text file\nWith multiple lines\nAnd some content."
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode='w') as f:
            f.write(txt_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert text == txt_content
        finally:
            os.unlink(fname)

    def test_read_json(self):
        """Test reading .json file."""
        json_content = '{"key": "value", "number": 42}'
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode='w') as f:
            f.write(json_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert "key" in text
            assert "value" in text
        finally:
            os.unlink(fname)

    def test_read_yaml(self):
        """Test reading .yaml file."""
        yaml_content = "key: value\nnumber: 42\nlist:\n  - item1\n  - item2"
        with tempfile.NamedTemporaryFile(suffix=".yaml", delete=False, mode='w') as f:
            f.write(yaml_content)
            f.flush()
            os.fsync(f.fileno())
            fname = f.name
        try:
            text = read_document(fname)
            assert "key" in text
            assert "value" in text
        finally:
            os.unlink(fname)

    def test_read_with_metadata(self):
        """Test reading with metadata extraction."""
        import docx
        doc = docx.Document()
        doc.add_heading("Metadata Test", 0)
        doc.add_paragraph("Content for metadata test.")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            try:
                text, meta = read_document_with_metadata(f.name)
                assert "Metadata Test" in text
                assert meta.format == ".docx"
                assert meta.size_bytes > 0
                assert meta.word_count > 0
                assert meta.char_count > 0
            finally:
                os.unlink(f.name)

    def test_file_not_found(self):
        """Test FileNotFoundError for missing file."""
        with pytest.raises(FileNotFoundError):
            read_document("/nonexistent/file.docx")

    def test_unsupported_format(self):
        """Test ValueError for unsupported format."""
        with tempfile.NamedTemporaryFile(suffix=".xyz", delete=False) as f:
            try:
                with pytest.raises(ValueError, match="Unsupported format"):
                    read_document(f.name)
            finally:
                os.unlink(f.name)

    def test_caching(self):
        """Test that caching works."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("Cache test content")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            try:
                reader = DocumentReader()
                # First read
                text1 = reader.read(f.name)
                # Second read (should use cache)
                text2 = reader.read(f.name)
                assert text1 == text2
            finally:
                os.unlink(f.name)

    def test_clear_cache(self):
        """Test cache clearing."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("Clear cache test")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            try:
                reader = DocumentReader()
                reader.read(f.name)
                assert len(reader._cache) == 1
                reader.clear_cache()
                assert len(reader._cache) == 0
            finally:
                os.unlink(f.name)


class TestDocumentMetadata:
    """Test DocumentMetadata dataclass."""

    def test_to_dict(self):
        """Test to_dict method."""
        meta = DocumentMetadata(
            path="/test/file.docx",
            format=".docx",
            size_bytes=1024,
            word_count=100,
            char_count=500,
            page_count=2,
            author="Test Author",
            created="2026-01-01T00:00:00",
            modified="2026-01-02T00:00:00",
            title="Test Title",
            extra={"custom": "value"}
        )
        d = meta.to_dict()
        assert d["path"] == "/test/file.docx"
        assert d["format"] == ".docx"
        assert d["size_bytes"] == 1024
        assert d["author"] == "Test Author"
        assert d["extra"]["custom"] == "value"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])